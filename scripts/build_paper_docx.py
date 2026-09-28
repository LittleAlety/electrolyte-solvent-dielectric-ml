"""Render the Chinese manuscript (``paper/paper_zh_draft.md``) to a ``.docx``.

The manuscript keeps a deliberately small Markdown vocabulary so that the
renderer stays readable: ATX headings, paragraphs, ``-`` bullets, pipe tables,
fenced code blocks, horizontal rules and images.  ``python-docx`` is not a
project dependency (the released artefacts are Markdown), so this script is run
with the bundled desktop runtime:

```
<bundled python> scripts/build_paper_docx.py --out "<path>.docx"
```
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPOSITORY_ROOT / "paper" / "paper_zh_draft.md"

BODY_FONT = "Times New Roman"
CJK_FONT = "\u7b49\u7ebf"  # 等线
MONO_FONT = "Consolas"
BODY_SIZE = Pt(10.5)
IMAGE_WIDTH_INCHES = 6.0

IMAGE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
BOLD = re.compile(r"\*\*(?P<text>[^*]+)\*\*")
ITALIC = re.compile(r"(?<!\*)\*(?!\*)(?P<text>[^*]+)\*(?!\*)")
TABLE_RULE = re.compile(r"^\|[\s:\-|]+\|$")


def _set_font(run, name: str = BODY_FONT, size: Pt = BODY_SIZE, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = size
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if bold is not None:
        run.bold = bold


def _add_runs(paragraph, text: str) -> None:
    """Add ``text`` to ``paragraph`` honouring ``**bold**`` and ``*italic*``."""

    tokens: list[tuple[str, bool, bool]] = []
    cursor = 0
    for match in BOLD.finditer(text):
        if match.start() > cursor:
            tokens.append((text[cursor : match.start()], False, False))
        tokens.append((match.group("text"), True, False))
        cursor = match.end()
    if cursor < len(text):
        tokens.append((text[cursor:], False, False))

    expanded: list[tuple[str, bool, bool]] = []
    for chunk, bold, _ in tokens:
        if bold:
            expanded.append((chunk, True, False))
            continue
        inner_cursor = 0
        for match in ITALIC.finditer(chunk):
            if match.start() > inner_cursor:
                expanded.append((chunk[inner_cursor : match.start()], False, False))
            expanded.append((match.group("text"), False, True))
            inner_cursor = match.end()
        if inner_cursor < len(chunk):
            expanded.append((chunk[inner_cursor:], False, False))

    if not expanded:
        expanded = [("", False, False)]
    for chunk, bold, italic in expanded:
        run = paragraph.add_run(chunk)
        _set_font(run, bold=bold)
        if italic:
            run.italic = True


def _add_heading(document: Document, text: str, level: int) -> None:
    sizes = {1: Pt(18), 2: Pt(14), 3: Pt(12)}
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _set_font(run, size=sizes[level], bold=True)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for position in range(width):
            value = row[position] if position < len(row) else ""
            paragraph = cells[position].paragraphs[0]
            _add_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _add_image(document: Document, src: str, alt: str) -> None:
    path = Path(src)
    if not path.is_absolute():
        path = REPOSITORY_ROOT / src
    if not path.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(paragraph, f"[missing figure: {src}]")
        return
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(path), width=Inches(IMAGE_WIDTH_INCHES))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    _set_font(run, size=Pt(9))
    run.italic = True
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def render(source: Path, out: Path) -> dict[str, object]:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    style = document.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    counts = {"headings": 0, "paragraphs": 0, "tables": 0, "images": 0, "bullets": 0, "code": 0}
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run("\n".join(block))
            run.font.name = MONO_FONT
            run.font.size = Pt(9)
            counts["code"] += 1
            continue

        image = IMAGE.match(stripped)
        if image:
            _add_image(document, image.group("src"), image.group("alt"))
            counts["images"] += 1
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and TABLE_RULE.match(lines[index + 1].strip()):
            rows = [_split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index].strip()))
                index += 1
            _add_table(document, rows)
            counts["tables"] += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            _add_heading(document, stripped[level:].strip(), min(level, 3))
            counts["headings"] += 1
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, stripped[2:].strip())
            counts["bullets"] += 1
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0.22)
        paragraph.paragraph_format.space_after = Pt(4)
        _add_runs(paragraph, stripped)
        counts["paragraphs"] += 1
        index += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    counts["bytes"] = out.stat().st_size
    counts["output"] = str(out)
    return counts


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    counts = render(args.source, args.out)
    for key in ("output", "bytes", "headings", "paragraphs", "bullets", "tables", "images", "code"):
        print(f"{key}: {counts[key]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
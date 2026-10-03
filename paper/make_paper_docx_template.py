'''Render the Chinese draft into the journal template layout.

The template (paper template 20233020) fixes the geometry and the type scale:
A4 with 2.3 cm side margins, body in Song / Times New Roman at 10.5 pt with a
fixed 15 pt line and 0.15 pt letter spacing, level-1 headings at 12 pt bold,
level-2/3 at 10.5 pt bold, figure captions at 9 pt bold, three-line tables with
7.5 pt text, and references at 7.5 pt.
'''

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

REPO = Path(__file__).resolve().parent.parent
MD = REPO / 'paper' / 'paper_zh_draft_v2.md'
OUT = REPO.parent / '成果输出' / '论文_电解液溶剂筛选中的描述符决策稳定性_期刊模版.docx'

SONG = '宋体'
HEI = '黑体'
FANG = '仿宋'
TNR = 'Times New Roman'
ARIAL = 'Arial'

BODY_SIZE = 10.5
SMALL_SIZE = 9.0
SIX_SIZE = 7.5
H1_SIZE = 12.0
TITLE_SIZE = 16.0
TITLE_EN_SIZE = 14.0
LINE_PT = 15.0
LETTER_SPACING = 3
TICK = chr(96)

TITLE_ZH = '电解液溶剂筛选中的描述符决策稳定性'
SUBTITLE_ZH = '——层级伪影、轴锁定分歧与筛查可用性：一个可审计的五通道漏斗'
AUTHORS_ZH = '作者姓名1*，作者姓名2'
AFFILIATION_ZH = '单位名称，城市 邮编'
CLC = 'O646；TP181'
KEYWORDS_ZH = '电解液溶剂筛选；描述符决策稳定性；电子结构层级；排序统计量；可审计机器学习；预注册；负结果登记'
TITLE_EN = ('Descriptor decision stability in electrolyte solvent screening: level artefacts, '
             'axis-locked disagreement and screening usability in an auditable five-channel funnel')
AUTHORS_EN = 'Author A, Author B'
AFFILIATION_EN = 'Affiliation, City Postcode, China'
KEYWORDS_EN = ('electrolyte solvent screening; descriptor decision stability; electronic-structure level; '
               'ranking statistics; auditable machine learning; pre-registration; negative results')
ACK = ('本工作的全部数字由仓库内可执行产物生成并经一致性脚本校验；负结果与封顶读数一并登记，未作删减。'
       '感谢开源社区提供的 ThermoML、NBS Circular 514、Batt-SLM 与 THEMol 等公开数据与工具。')

ABSTRACT_ZH = (
    '本文研究电解液溶剂筛选中的描述符决策稳定性：当一条描述符的排序本身不稳定时，任何基于它的筛选结论都不成立。'
    '我们构建并公开 246 种近室温纯有机液体的静态介电常数数据集（v0.3.3，38 字段，按来源分层收录），'
    '并另建轨道与氧化还原通道（HOMO/LUMO/gap 与氧化还原电位）。方法上提出单变量台阶框架：'
    '自由分子电子结构轴（P0 Koopmans 轨道 → P1 ΔSCF → P2 隐式溶剂）与条件态轴'
    '（C0 裸分子 → C1 [Li(M)]+ → C2 [Li(M)2]+），每级只改一个因素；判据跑前冻结，'
    '靶值置换必须使读数塌缩，同一化合物不得跨折，失败写入负结果登记。三条主结果：'
    '决策稳定性可被量化，排序统计量 tau_b 的重抽标准差服从有限总体修正的闭式，给出小样本档的噪声地板 0.122；'
    '还原轴上的跨工作分歧是层级伪影，氧化轴六档全部不超过 0.04 sigma，而还原轴全部不小于 2.3 sigma；'
    '排序统计量与筛查可用性脱钩，tau_b 与 Top-10 重叠率的秩相关仅 +0.793，未解析占比与它的秩相关是 -0.862。'
    '在化合物分组口径下，介电通道的冻结基线为 0.4091、冻结头条为 0.4766；表示侧增广杠杆已被逐一实测证否，'
    '开放许可化合物名册实测封顶，因此 0.70 在现有约束下没有路径。')

ABSTRACT_EN = (
    'When the ordering induced by a descriptor is itself unstable, every screening conclusion built on '
    'it is unsound. We study that problem for electrolyte solvent screening and release an auditable, '
    'machine-learning-ready dataset of 246 near-room-temperature pure organic liquids (static '
    'permittivity, 38 fields, stratified by source), together with separately built orbital and redox '
    'channels. The method is a single-variable ladder: a free-molecule electronic-structure axis '
    '(P0 Koopmans orbital, P1 delta-SCF, P2 implicit solvent) and a conditional-state axis (C0 bare '
    'molecule, C1 [Li(M)]+, C2 [Li(M)2]+), one factor per rung. Criteria are frozen before the run, '
    'target permutation must collapse the reading, no compound may straddle a fold, and every failure '
    'is written into the delivery. Three results follow. First, decision stability is measurable: the '
    'resampling standard deviation of the ranking statistic tau_b follows a finite-population '
    'correction, which puts a noise floor of 0.122 on small-sample comparisons. Second, the '
    'cross-study disagreement on the reduction axis is a level artefact: all six oxidation-axis '
    'rungs sit within 0.04 sigma while all three reduction-axis rungs sit beyond 2.3 sigma, and '
    'holding the compounds fixed while changing only the level flips tau_b from +0.835 to -0.602. '
    'Third, ranking statistics and screening usability come apart: tau_b correlates with Top-10 '
    'overlap at only +0.793, while the unresolved fraction correlates with it at -0.862. Under '
    'grouped compound splits the frozen baseline and headline of the permittivity channel are '
    '0.4091 and 0.4766, and the viscosity channel misses its 0.15 gate by 0.025. Every '
    'representation-side lever has been measured and refuted, and the open-licence compound roster '
    'is capped, so 0.70 has no path under current constraints. We report the negative results, the '
    'ceiling and the audit trail in full.')

def letter_space(run, twentieths=LETTER_SPACING):
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), str(twentieths))
    rpr.append(spacing)


def font(run, ascii_font, cjk, size, bold=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), cjk)
    rfonts.set(qn('w:ascii'), ascii_font)
    rfonts.set(qn('w:hAnsi'), ascii_font)
    letter_space(run)


def style_paragraph(par, *, line_pt=LINE_PT, before=0.0, after=0.0, indent_chars=0.0, align=None, hanging=None, keep_with_next=False):
    fmt = par.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if indent_chars:
        fmt.first_line_indent = Pt(indent_chars * BODY_SIZE)
    if hanging is not None:
        fmt.left_indent = Pt(hanging)
        fmt.first_line_indent = Pt(-hanging)
    if align is not None:
        par.alignment = align
    fmt.keep_with_next = keep_with_next
    return par


def add_text(par, text, size, ascii_font, cjk, bold_default=False):
    text = text.replace(TICK, '')
    for index, piece in enumerate(text.split('**')):
        if not piece:
            continue
        run = par.add_run(piece)
        if index % 2 == 1:
            font(run, ascii_font, cjk, size, bold=True)
        else:
            font(run, ascii_font, cjk, size, bold=True if bold_default else None)
    return par


def add_paragraph(doc, text, *, size=BODY_SIZE, ascii_font=TNR, cjk=SONG, indent=True, align=None,
                  before=0.0, after=0.0, bold_default=False, hanging=None, line_pt=LINE_PT):
    par = doc.add_paragraph()
    style_paragraph(par, line_pt=line_pt, before=before, after=after,
                    indent_chars=2.0 if indent else 0.0, align=align, hanging=hanging)
    add_text(par, text, size, ascii_font, cjk, bold_default=bold_default)
    return par


def cell_border(cell, edge, size_eighths):
    tag = 'w:' + edge
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcpr.append(borders)
    element = borders.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        borders.append(element)
    element.set(qn('w:val'), 'single')
    element.set(qn('w:sz'), str(size_eighths))
    element.set(qn('w:space'), '0')
    element.set(qn('w:color'), '000000')


def add_three_line_table(doc, rows):
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column in range(width):
            cell = cells[column]
            cell.text = ''
            par = cell.paragraphs[0]
            style_paragraph(par, line_pt=10.0,
                            align=WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else None)
            add_text(par, row[column] if column < len(row) else '', SIX_SIZE, TNR, SONG,
                     bold_default=(row_index == 0))
    for cell in table.rows[0].cells:
        cell_border(cell, 'top', 12)
        cell_border(cell, 'bottom', 6)
    for cell in table.rows[-1].cells:
        cell_border(cell, 'bottom', 12)
    doc.add_paragraph()


def parse_image(line):
    if not line.startswith('!['):
        return None
    body = line[2:]
    marker = body.find('](')
    if marker < 0 or not body.endswith(')'):
        return None
    return body[:marker], body[marker + 2:-1]


def add_figure(doc, relative_path, caption):
    target = Path(relative_path)
    if not target.is_absolute():
        target = REPO / relative_path
    if not target.is_file():
        add_paragraph(doc, '[缺图：' + relative_path + ']', size=SMALL_SIZE, indent=False)
        return False
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, before=4.0, after=2.0, keep_with_next=True)
    par.add_run().add_picture(str(target), width=Cm(15.0))
    cap = doc.add_paragraph()
    style_paragraph(cap, align=WD_ALIGN_PARAGRAPH.CENTER, before=0.0, after=6.0)
    add_text(cap, caption, SMALL_SIZE, ARIAL, HEI, bold_default=True)
    return True

def build_front_matter(doc):
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, before=0.0, after=2.0)
    add_text(par, TITLE_ZH, TITLE_SIZE, ARIAL, SONG, bold_default=True)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, after=6.0)
    add_text(par, SUBTITLE_ZH, H1_SIZE, TNR, SONG, bold_default=True)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, after=0.0)
    add_text(par, AUTHORS_ZH, H1_SIZE, TNR, SONG)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, after=6.0)
    add_text(par, AFFILIATION_ZH, SMALL_SIZE, TNR, SONG)
    add_paragraph(doc, '**摘要：**' + ABSTRACT_ZH, size=SMALL_SIZE, indent=False, before=6.0)
    add_paragraph(doc, '**关键词：**' + KEYWORDS_ZH, size=SMALL_SIZE, indent=False, after=0.0)
    add_paragraph(doc, '**中图分类号：**' + CLC, size=SMALL_SIZE, indent=False, after=8.0)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, before=6.0, after=2.0)
    add_text(par, TITLE_EN, TITLE_EN_SIZE, ARIAL, ARIAL, bold_default=True)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, after=0.0)
    add_text(par, AUTHORS_EN, H1_SIZE, ARIAL, ARIAL)
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, after=6.0)
    add_text(par, AFFILIATION_EN, SMALL_SIZE, TNR, ARIAL)
    add_paragraph(doc, '**Abstract:** ' + ABSTRACT_EN, size=SMALL_SIZE, ascii_font=ARIAL, cjk=ARIAL, indent=False, before=6.0)
    add_paragraph(doc, '**Key Words:** ' + KEYWORDS_EN, size=SMALL_SIZE, ascii_font=ARIAL, cjk=ARIAL, indent=False, after=10.0)


def is_table_separator(line):
    stripped = line.strip()
    return bool(stripped) and set(stripped) <= set('-:| ') and '-' in stripped


def split_table_row(line):
    stripped = line.strip()
    if stripped.startswith('|'):
        stripped = stripped[1:]
    if stripped.endswith('|'):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split('|')]


def is_reference_entry(line):
    stripped = line.strip()
    head = stripped.split('.')[0] if '.' in stripped else ''
    return bool(head) and head.isdigit() and stripped.startswith(head + '. ')


def add_heading(doc, text, level):
    size = H1_SIZE if level == 1 else BODY_SIZE
    par = doc.add_paragraph()
    style_paragraph(par, before=8.0 if level == 1 else 5.0, after=3.0, keep_with_next=True)
    add_text(par, text, size, ARIAL, HEI, bold_default=True)
    return par


def render_body(doc, lines, start, stop):
    index = start
    figures = 0
    tables = 0
    while index < stop:
        line = lines[index]
        stripped = line.strip()
        if not stripped or stripped == '---':
            index += 1
            continue
        image = parse_image(stripped)
        if image is not None:
            if add_figure(doc, image[1], image[0]):
                figures += 1
            index += 1
            continue
        if stripped.startswith('|'):
            rows = []
            while index < stop and lines[index].strip().startswith('|'):
                if not is_table_separator(lines[index]):
                    rows.append(split_table_row(lines[index]))
                index += 1
            if rows:
                add_three_line_table(doc, rows)
                tables += 1
            continue
        if stripped.startswith('#### '):
            add_heading(doc, stripped[5:], 3)
        elif stripped.startswith('### '):
            add_heading(doc, stripped[4:], 2)
        elif stripped.startswith('## '):
            add_heading(doc, stripped[3:], 1)
        elif stripped.startswith('# '):
            add_heading(doc, stripped[2:], 1)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            add_paragraph(doc, '· ' + stripped[2:], indent=False, after=0.0)
        else:
            add_paragraph(doc, stripped, after=0.0)
        index += 1
    return figures, tables

def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--md', default=str(MD))
    parser.add_argument('--out', default=str(OUT))
    return parser.parse_args()


def center_heading(doc, text):
    par = doc.add_paragraph()
    style_paragraph(par, align=WD_ALIGN_PARAGRAPH.CENTER, before=10.0, after=4.0, keep_with_next=True)
    add_text(par, text, SMALL_SIZE, ARIAL, HEI, bold_default=True)
    return par


def main():
    args = parse_args()
    lines = Path(args.md).read_text(encoding='utf-8').split(chr(10))
    abstract_index = None
    body_index = None
    reference_index = None
    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped == '## 摘要' and abstract_index is None:
            abstract_index = index
        if stripped.startswith('## 1 引言') and body_index is None:
            body_index = index
        if stripped == '## 参考文献' and reference_index is None:
            reference_index = index
    if body_index is None or reference_index is None:
        raise SystemExit('the draft is missing its body or its reference section')
    structured = []
    if abstract_index is not None:
        for line in lines[abstract_index + 1:body_index]:
            stripped = line.strip()
            if not stripped or stripped == '---' or stripped.startswith('**关键词**'):
                continue
            structured.append(stripped)
    metadata = []
    for line in lines[:body_index]:
        stripped = line.strip()
        if stripped.startswith('**数据版本**') or stripped.startswith('**代码与交付版本**'):
            metadata.append(stripped)
        elif stripped.startswith('**永久归档**') or stripped.startswith('**许可**') or stripped.startswith('**口径纪律**'):
            metadata.append(stripped)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.top_margin = Cm(3.4)
    section.bottom_margin = Cm(3.0)
    build_front_matter(doc)
    figures, tables = render_body(doc, lines, body_index, reference_index)
    par = doc.add_paragraph()
    style_paragraph(par, before=10.0, after=10.0)
    add_text(par, '**致谢：**' + ACK, SMALL_SIZE, TNR, FANG)
    center_heading(doc, '参  考  文  献')
    for line in lines[reference_index + 1:]:
        stripped = line.strip()
        if not stripped or stripped == '---':
            continue
        if stripped.startswith('*本文全部数字'):
            continue
        if is_reference_entry(stripped):
            add_paragraph(doc, stripped, size=SIX_SIZE, indent=False, hanging=12.0, line_pt=11.0)
        else:
            add_paragraph(doc, stripped, size=SIX_SIZE, indent=False, after=2.0, line_pt=11.0)
    center_heading(doc, '版本与许可声明')
    for line in metadata:
        add_paragraph(doc, line, size=SIX_SIZE, indent=False, after=1.0, line_pt=11.0)
    add_heading(doc, '附录 D　结构化摘要（详版）', 1)
    for line in structured:
        add_paragraph(doc, line, size=SMALL_SIZE, after=1.0)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print('saved ' + str(out))
    print('paragraphs ' + str(len(doc.paragraphs)) + ' | tables ' + str(tables) + ' | figures ' + str(figures))


if __name__ == '__main__':
    main()

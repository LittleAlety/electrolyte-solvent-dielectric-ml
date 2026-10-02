# -*- coding: utf-8 -*-
"""Render paper/paper_zh_draft_v2.md into a Word document with headings, tables and figures."""
import argparse
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent
MD = REPO / "paper" / "paper_zh_draft_v2.md"
DEFAULT_OUT = REPO.parent / "\u6210\u679c\u8f93\u51fa" / "\u8bba\u6587_\u7535\u89e3\u6db2\u6eb6\u5242\u7b5b\u9009\u4e2d\u7684\u63cf\u8ff0\u7b26\u51b3\u7b56\u7a33\u5b9a\u6027.docx"

IMAGE_RE = re.compile(r"^!\[(?P<cap>.*?)\]\((?P<path>.*?)\)\s*$")
BOLD_RE = re.compile(r"(\*\*.+?\*\*|\x60[^\x60]+\x60)")


def set_cjk(style_name, doc, ascii_font, cjk_font, size=None, bold=None):
    st = doc.styles[style_name]
    st.font.name = ascii_font
    if size is not None:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), cjk_font)
    rf.set(qn("w:ascii"), ascii_font)
    rf.set(qn("w:hAnsi"), ascii_font)


def add_runs(par, text):
    for piece in BOLD_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            r.bold = True
        elif piece.startswith("\x60") and piece.endswith("\x60") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        else:
            par.add_run(piece)


def add_image(doc, rel_path, caption):
    img = REPO / rel_path
    if not img.exists():
        par = doc.add_paragraph()
        par.add_run("[\u7f3a\u56fe\uff1a" + rel_path + "]").italic = True
        return False
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width = Inches(6.1)
    par.add_run().add_picture(str(img), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return True


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = tbl.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cells[j].text = ""
            par = cells[j].paragraphs[0]
            add_runs(par, txt)
            for r in par.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def is_sep(line):
    s = line.strip()
    return bool(s) and set(s) <= set("|-: ") and "-" in s


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", default=str(MD))
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    args = ap.parse_args()
    lines = io.open(args.md, encoding="utf-8").read().split("\n")

    doc = Document()
    set_cjk("Normal", doc, "Times New Roman", "\u5b8b\u4f53", size=10.5)
    for name, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("Title", 20)):
        try:
            set_cjk(name, doc, "Times New Roman", "\u9ed1\u4f53", size=sz)
        except KeyError:
            pass
    sec = doc.sections[0]
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    n_fig = 0
    n_tab = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s == "":
            i += 1
            continue
        if s == "---":
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for bl in block:
                par = doc.add_paragraph()
                r = par.add_run(bl if bl else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                par.paragraph_format.space_after = Pt(0)
            continue
        m = IMAGE_RE.match(s)
        if m:
            if add_image(doc, m.group("path"), m.group("cap")):
                n_fig += 1
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
                n_tab += 1
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=2)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=1)
        elif s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=0)
        else:
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Inches(0.28)
            par.paragraph_format.space_after = Pt(4)
            add_runs(par, s)
        i += 1

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print("saved", out)
    print("paragraphs", len(doc.paragraphs), "tables", n_tab, "figures", n_fig)


if __name__ == "__main__":
    main()
